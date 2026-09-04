"""
File-format templates ("configs" in v2).

A template is the small two-column workbook v2 used to describe where the data lives in a
machine's output file and what to do to it. Three shapes exist:

  * two-channel  (16 rows) — IDEAL-CT, SCB, HT-IDT, Bond Strength
  * four-channel (28 rows) — Cold SCB, DCT   (time / stroke / load / COD)
  * beam         (12 rows) — Beam Fatigue     (cycles / flexural stiffness / E)

Row order is significant: v2 reads the template by row *index*, not by label, and the port
does the same so that the lab's existing template library imports unchanged.
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Dict, List, Optional

import openpyxl

KIND_XY = "xy"          # two-channel
KIND_FOUR = "four"      # four-channel
KIND_BEAM = "beam"


@dataclass
class ChannelTransform:
    offset_start: int = 0   # subtract first value
    negative: int = 0       # flip sign
    offset_end: int = 0     # subtract last value

    def as_tuple(self):
        return (int(self.offset_start), int(self.negative), int(self.offset_end))


@dataclass
class Template:
    id: str
    name: str
    kind: str
    columns: Dict[str, int]                 # channel -> 1-based column index
    data_start: int = 1
    data_end: Optional[int] = None          # None == Inf
    delimiter: str = ","
    sheet: int = 1                          # 1-based
    transforms: Dict[str, ChannelTransform] = field(default_factory=dict)
    units: Dict[str, str] = field(default_factory=dict)
    builtin: bool = False
    source: str = ""                        # original filename, for reference

    # -- (de)serialisation -------------------------------------------------
    def to_dict(self) -> dict:
        d = asdict(self)
        return d

    @staticmethod
    def from_dict(d: dict) -> "Template":
        tr = {k: ChannelTransform(**v) for k, v in d.get("transforms", {}).items()}
        return Template(id=d["id"], name=d["name"], kind=d["kind"], columns=dict(d["columns"]),
                        data_start=int(d.get("data_start", 1)), data_end=d.get("data_end"),
                        delimiter=d.get("delimiter", ","), sheet=int(d.get("sheet", 1)),
                        transforms=tr, units=dict(d.get("units", {})),
                        builtin=bool(d.get("builtin", False)), source=d.get("source", ""))


# ---------------------------------------------------------------------------
# Parsing the v2 workbook format
# ---------------------------------------------------------------------------

def _rows(path: Path) -> List[List]:
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    out = []
    for r in ws.iter_rows(values_only=True):
        out.append(list(r))
    return out


def _num(v, default=None):
    if v is None or v == "":
        return default
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    if s.lower() in ("inf", "infinity"):
        return None
    try:
        return float(s)
    except ValueError:
        return default


def _int(v, default):
    n = _num(v, None)
    return int(n) if n is not None else default


def _str(v, default=""):
    return default if v is None else str(v)


def parse_template_workbook(path: Path, template_id: Optional[str] = None,
                            name: Optional[str] = None, builtin: bool = False) -> Template:
    """Parse a v2-format template workbook. The shape is inferred from the row count /
    first-column labels."""
    rows = _rows(path)
    val = lambda i: rows[i - 1][1] if i - 1 < len(rows) and len(rows[i - 1]) > 1 else None
    labels = [(_str(r[0]) if r else "") for r in rows]
    n = len([l for l in labels if l])
    tid = template_id or re.sub(r"[^A-Za-z0-9_.-]+", "_", path.stem)
    nm = name or path.stem

    if n >= 28 and labels[0].lower().startswith("time"):
        # four-channel: rows 1-4 columns, 5 start, 6 end, 7-10 names, 11 delim,
        # 12-23 transforms (see readers.read_four_channel for the v2 mapping), 24-27 units, 28 sheet
        cols = {"time": _int(val(1), 1), "stroke": _int(val(2), 2),
                "load": _int(val(3), 3), "cod": _int(val(4), 4)}
        tr = {
            "row12": ChannelTransform(_int(val(12), 0), _int(val(13), 0), _int(val(14), 0)),
            "row15": ChannelTransform(_int(val(15), 0), _int(val(16), 0), _int(val(17), 0)),
            "row18": ChannelTransform(_int(val(18), 0), _int(val(19), 0), _int(val(20), 0)),
            "row21": ChannelTransform(_int(val(21), 0), _int(val(22), 0), _int(val(23), 0)),
        }
        units = {"time": _str(val(24), "sec"), "stroke": _str(val(25), "mm"),
                 "load": _str(val(26), "kN"), "cod": _str(val(27), "mm")}
        return Template(tid, nm, KIND_FOUR, cols, _int(val(5), 1), _num(val(6)),
                        _str(val(11), ","), _int(val(28), 1), tr, units, builtin, path.name)

    if n >= 12 and labels[0].lower().startswith("cycles"):
        cols = {"cycles": _int(val(1), 1), "stiffness": _int(val(2), 2), "modulus": _int(val(3), 3)}
        units = {"stiffness": _str(val(10), "MPa"), "modulus": _str(val(11), "MPa")}
        return Template(tid, nm, KIND_BEAM, cols, _int(val(4), 1), _num(val(5)),
                        _str(val(9), ","), _int(val(12), 1), {}, units, builtin, path.name)

    # two-channel (15 or 16 rows; row 16 = sheet may be absent)
    cols = {"x": _int(val(1), 1), "y": _int(val(2), 2)}
    tr = {"x": ChannelTransform(_int(val(8), 0), _int(val(9), 0), _int(val(10), 0)),
          "y": ChannelTransform(_int(val(11), 0), _int(val(12), 0), _int(val(13), 0))}
    units = {"x": _str(val(14), "mm"), "y": _str(val(15), "kN")}
    return Template(tid, nm, KIND_XY, cols, _int(val(3), 1), _num(val(4)),
                    _str(val(7), ","), _int(val(16), 1), tr, units, builtin, path.name)


def write_template_workbook(t: Template, path: Path) -> None:
    """Write a template back out in the v2 workbook format (so v2 users can share)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    end = "Inf" if t.data_end is None else t.data_end
    if t.kind == KIND_FOUR:
        r = t.transforms
        rows = [("Time Start", t.columns["time"]), ("Stroke Start", t.columns["stroke"]),
                ("Load Start", t.columns["load"]), ("COD Gage Start", t.columns["cod"]),
                ("Data Start", t.data_start), ("Data End", end), ("X1", "Time"), ("X2", "Stroke"),
                ("X3", "Load"), ("X4", "COD"), ("Delimiter", t.delimiter)]
        for key, lab in (("row12", "Time"), ("row15", "Stroke"), ("row18", "COD Gage"), ("row21", "Load")):
            c = r.get(key, ChannelTransform())
            rows += [(f"Offset from Start ({lab})", c.offset_start), (f"Negative ({lab})", c.negative),
                     (f"Offset from End ({lab})", c.offset_end)]
        rows += [("Time units", t.units.get("time", "sec")), ("Stroke Units", t.units.get("stroke", "mm")),
                 ("Load Units", t.units.get("load", "kN")), ("COD Gage units", t.units.get("cod", "mm")),
                 ("Sheet", t.sheet)]
    elif t.kind == KIND_BEAM:
        rows = [("Cycles", t.columns["cycles"]), ("Flexural Stiffness", t.columns["stiffness"]),
                ("Elastic Modulus", t.columns["modulus"]), ("Data Start", t.data_start), ("Data End", end),
                ("X", "Cycles"), ("Y1", "Flexural Stiffness"), ("Y2", "Elastic Modulus"),
                ("Delimiter", t.delimiter), ("Flexural Stiffness Units", t.units.get("stiffness", "MPa")),
                ("E Units", t.units.get("modulus", "MPa")), ("Sheet", t.sheet)]
    else:
        x = t.transforms.get("x", ChannelTransform())
        y = t.transforms.get("y", ChannelTransform())
        rows = [("Displacement", t.columns["x"]), ("Load Start", t.columns["y"]), ("Data Start", t.data_start),
                ("Data End", end), ("X", "Displacement"), ("Y", "Load"), ("Delimiter", t.delimiter),
                ("Offset from Start (Disp)", x.offset_start), ("Negative (Disp)", x.negative),
                ("Offset from End (Disp)", x.offset_end), ("Offset from Start (Load)", y.offset_start),
                ("Negative (Load)", y.negative), ("Offset from End (Load)", y.offset_end),
                ("Displacement units", t.units.get("x", "mm")), ("Load Units", t.units.get("y", "kN")),
                ("Sheet", t.sheet)]
    for a, b in rows:
        ws.append([a, b])
    wb.save(str(path))


# ---------------------------------------------------------------------------
# Library: bundled defaults + user templates on disk
# ---------------------------------------------------------------------------

MODULE_KIND = {
    "ideal_ct": KIND_XY, "scb": KIND_XY, "htidt": KIND_XY, "bond": KIND_XY,
    "cold_scb": KIND_FOUR, "dct": KIND_FOUR, "beam": KIND_BEAM,
}


class TemplateLibrary:
    def __init__(self, builtin_dir: Path, user_dir: Path):
        self.builtin_dir = Path(builtin_dir)
        self.user_dir = Path(user_dir)
        self.user_dir.mkdir(parents=True, exist_ok=True)
        self._items: Dict[str, Template] = {}
        self.reload()

    def reload(self) -> None:
        self._items.clear()
        for p in sorted(self.builtin_dir.glob("*.xlsx")):
            try:
                t = parse_template_workbook(p, builtin=True)
                self._items[t.id] = t
            except Exception:
                continue
        for p in sorted(self.user_dir.glob("*.json")):
            try:
                t = Template.from_dict(json.loads(p.read_text()))
                self._items[t.id] = t
            except Exception:
                continue

    def list(self, kind: Optional[str] = None) -> List[Template]:
        items = list(self._items.values())
        if kind:
            items = [t for t in items if t.kind == kind]
        return sorted(items, key=lambda t: (not t.builtin, t.name.lower()))

    def get(self, tid: str) -> Template:
        return self._items[tid]

    def save(self, t: Template) -> Template:
        t.builtin = False
        (self.user_dir / f"{t.id}.json").write_text(json.dumps(t.to_dict(), indent=2))
        self._items[t.id] = t
        return t

    def delete(self, tid: str) -> None:
        t = self._items.get(tid)
        if t and not t.builtin:
            (self.user_dir / f"{tid}.json").unlink(missing_ok=True)
            del self._items[tid]

    def import_workbook(self, path: Path, name: Optional[str] = None) -> Template:
        t = parse_template_workbook(Path(path), name=name)
        return self.save(t)
